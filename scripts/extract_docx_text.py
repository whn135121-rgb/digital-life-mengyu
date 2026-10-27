from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_block_items(parent):
    if isinstance(parent, DocumentObject):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)


def table_rows(table: Table):
    for row in table.rows:
        cells = []
        for cell in row.cells:
            parts = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
            for nested in cell.tables:
                for nested_row in nested.rows:
                    parts.append(" ".join(c.text.strip() for c in nested_row.cells if c.text.strip()))
            cells.append(" / ".join(parts))
        if any(cells):
            yield cells


def main() -> None:
    src = Path(sys.argv[1])
    out = Path(sys.argv[2])
    doc = Document(str(src))
    lines: list[str] = []
    lines.append(f"FILE: {src}")
    cp = doc.core_properties
    lines.append(f"TITLE: {cp.title or ''}")
    lines.append(f"AUTHOR: {cp.author or ''}")
    lines.append(f"CREATED: {cp.created}")
    lines.append(f"MODIFIED: {cp.modified}")
    lines.append(f"LAST_MODIFIED_BY: {cp.last_modified_by or ''}")
    lines.append(f"REVISION: {cp.revision}")
    lines.append(f"PARAGRAPHS: {len(doc.paragraphs)}  TABLES: {len(doc.tables)}")
    lines.append("")
    lines.append("=== BODY ===")
    table_no = 0
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            style = block.style.name if block.style else ""
            if not text:
                continue
            if style.startswith("Heading") or style in {"Title", "Subtitle"}:
                lines.append(f"\n[{style}] {text}")
            else:
                lines.append(f"[{style}] {text}")
        else:
            table_no += 1
            lines.append(f"\n--- TABLE {table_no} ({len(block.rows)} rows x {len(block.columns)} cols) ---")
            for i, row in enumerate(table_rows(block), 1):
                lines.append(f"R{i}: " + " | ".join(row))
            lines.append("--- END TABLE ---")
    lines.append("")
    lines.append("=== SECTIONS ===")
    for i, section in enumerate(doc.sections, 1):
        lines.append(f"SECTION {i}: page {section.page_width}x{section.page_height}, margins L{section.left_margin} R{section.right_margin} T{section.top_margin} B{section.bottom_margin}")
        for name, header_footer in (("HEADER", section.header), ("FOOTER", section.footer)):
            txt = " / ".join(p.text.strip() for p in header_footer.paragraphs if p.text.strip())
            if txt:
                lines.append(f"  {name}: {txt}")
    out.write_text("\n".join(lines), encoding="utf-8")
    print(f"WROTE {out} ({len(lines)} lines)")


if __name__ == "__main__":
    main()
